#!/usr/bin/env python3
"""Generate deterministic openpyxl-authored XLSX import fixtures."""

from __future__ import annotations

import argparse
import re
from datetime import datetime, timedelta, timezone
from pathlib import Path
from tempfile import NamedTemporaryFile
from zipfile import ZIP_DEFLATED, ZipFile, ZipInfo

import openpyxl
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[1]
OUTPUT_DIRECTORY = ROOT / "src" / "test" / "fixtures" / "xlsx"
FIXED_DATETIME = datetime(2026, 1, 1, tzinfo=timezone.utc)
FIXED_ZIP_TIME = (2026, 1, 1, 0, 0, 0)
EXPECTED_OPENPYXL_VERSION = "3.1.5"


def base_workbook() -> tuple[Workbook, object]:
    workbook = new_workbook()
    worksheet = workbook.active
    worksheet.title = "S"
    worksheet.append(["Q", "V"])
    for quarter, value in (("Q1", 120), ("Q2", 150), ("Q3", 90), ("Q4", 200)):
        worksheet.append([quarter, value])
    return workbook, worksheet


def new_workbook() -> Workbook:
    workbook = Workbook()
    workbook.properties.created = FIXED_DATETIME
    workbook.properties.modified = FIXED_DATETIME
    return workbook


def save_deterministic(workbook: Workbook, destination: str | Path) -> None:
    destination = OUTPUT_DIRECTORY / destination if isinstance(destination, str) else destination
    destination.parent.mkdir(parents=True, exist_ok=True)
    with NamedTemporaryFile(suffix=".xlsx", dir="/tmp") as temporary:
        workbook.save(temporary.name)
        with ZipFile(temporary.name, "r") as source, ZipFile(
            destination, "w", compression=ZIP_DEFLATED, compresslevel=9
        ) as output:
            for source_info in source.infolist():
                target_info = ZipInfo(source_info.filename, FIXED_ZIP_TIME)
                target_info.compress_type = ZIP_DEFLATED
                target_info.external_attr = source_info.external_attr
                target_info.create_system = source_info.create_system
                content = source.read(source_info.filename)
                if source_info.filename == "docProps/core.xml":
                    content = re.sub(
                        rb"(<dcterms:modified\b[^>]*>)[^<]*(</dcterms:modified>)",
                        rb"\g<1>2026-01-01T00:00:00Z\g<2>",
                        content,
                    )
                output.writestr(target_info, content)


def generate_chart_fixture() -> None:
    workbook, worksheet = base_workbook()
    chart = BarChart()
    chart.add_data(Reference(worksheet, min_col=2, min_row=1, max_row=5), titles_from_data=True)
    chart.set_categories(Reference(worksheet, min_col=1, min_row=2, max_row=5))
    worksheet.add_chart(chart, "D2")
    save_deterministic(workbook, "variant-chart.xlsx")


def generate_comment_fixture() -> None:
    workbook, worksheet = base_workbook()
    worksheet["A1"].comment = Comment("hello", "author")
    save_deterministic(workbook, "variant-comment.xlsx")


def generate_table_fixture() -> None:
    workbook, worksheet = base_workbook()
    table = Table(displayName="T1", ref="A1:B5")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    worksheet.add_table(table)
    save_deterministic(workbook, "variant-table.xlsx")


def generate_kitchen_sink_fixture(
    dense_rows: int = 250,
    destination: str | Path = "real-kitchen-sink-trimmed.xlsx",
) -> None:
    workbook = new_workbook()
    sales = workbook.active
    sales.title = "Sales"
    sales.freeze_panes = "A2"
    sales.append(["Date", "Rep", "Region", "Units", "Unit Price", "Amount"])
    regions = ("North", "South", "East", "West")
    start = datetime(2026, 2, 1)
    for index in range(1, 51):
        row = index + 1
        sales.append([
            start + timedelta(days=index),
            f"Rep {(index - 1) % 5 + 1}",
            regions[index % len(regions)],
            index * 2,
            9.99 + (index % 7) * 5,
            f"=D{row}*E{row}",
        ])
        sales.cell(row, 1).number_format = "yyyy-mm-dd"

    table = Table(displayName="SalesTable", ref="A1:F51")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium9", showRowStripes=True)
    sales.add_table(table)
    chart = BarChart()
    chart.add_data(Reference(sales, min_col=6, min_row=1, max_row=11), titles_from_data=True)
    chart.set_categories(Reference(sales, min_col=2, min_row=2, max_row=11))
    sales.add_chart(chart, "H2")

    dashboard = workbook.create_sheet("Dashboard")
    dashboard.merge_cells("A1:D1")
    dashboard["A1"] = "Sales dashboard"
    dashboard["A9"] = "Docs"
    dashboard["A9"].hyperlink = "https://example.com/report"
    dashboard["A10"] = "Reviewed"
    dashboard["A10"].comment = Comment("Checked by finance", "Finance")

    dense = workbook.create_sheet("Data10k")
    for _ in range(dense_rows):
        dense.append([1] * 8)

    save_deterministic(workbook, destination)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--full-kitchen-sink",
        type=Path,
        metavar="OUTPUT",
        help="generate only the untrimmed 80,000-cell kitchen sink at OUTPUT",
    )
    arguments = parser.parse_args()
    if openpyxl.__version__ != EXPECTED_OPENPYXL_VERSION:
        raise RuntimeError(
            f"Expected openpyxl {EXPECTED_OPENPYXL_VERSION}, found {openpyxl.__version__}"
        )
    if arguments.full_kitchen_sink:
        generate_kitchen_sink_fixture(10_000, arguments.full_kitchen_sink.resolve())
        return

    generate_chart_fixture()
    generate_comment_fixture()
    generate_table_fixture()
    generate_kitchen_sink_fixture()


if __name__ == "__main__":
    main()
